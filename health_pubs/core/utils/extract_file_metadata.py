# core/utils/extract_file_metadata.py
from __future__ import annotations

import json
import logging
import mimetypes
import os
import shutil
import subprocess
import tempfile
import time
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Union
from urllib.parse import urlparse, unquote

import boto3
from botocore.exceptions import ClientError
from django.conf import settings
from django.core.cache import cache

from pptx import Presentation
from docx import Document as DocxDocument
from PyPDF2 import PdfReader
from PIL import ImageFile  # incremental image parser

logger = logging.getLogger(__name__)
if not logger.handlers:
    logging.basicConfig(
        level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s"
    )

DEFAULT_MIME = "application/octet-stream"
s3_client = boto3.client("s3")

# -------------------- configuration --------------------
STRICT_DOC_PAGE_META = bool(getattr(settings, "STRICT_DOC_PAGE_META", True))
DOC_PAGECOUNT_VIA_LIBREOFFICE = bool(
    getattr(settings, "DOC_PAGECOUNT_VIA_LIBREOFFICE", True)
)
DOCX_INCLUDE_PAGECOUNT = bool(
    getattr(settings, "DOCX_INCLUDE_PAGECOUNT", True)
)  # ENABLED

# Resolve the binary dynamically; default to the *name* so we can find it with shutil.which
LIBREOFFICE_BIN = getattr(settings, "LIBREOFFICE_BIN", "soffice")
LIBREOFFICE_TIMEOUT_SECS = int(getattr(settings, "LIBREOFFICE_TIMEOUT_SECS", 30))

FILE_METADATA_CACHE_TTL = int(getattr(settings, "FILE_METADATA_CACHE_TTL", 6 * 60 * 60))
FILE_METADATA_TIME_BUDGET_MS = int(
    getattr(settings, "FILE_METADATA_TIME_BUDGET_MS", 300) or 0
)
MAX_METADATA_BYTES = getattr(settings, "MAX_METADATA_BYTES", 2 * 1024 * 1024)

DOC_DEEP_PROBE_EXTS = {
    e.lower()
    for e in getattr(
        settings, "DOC_DEEP_PROBE_EXTS", ["pdf", "pptx", "docx", "doc", "odt", "ppt"]
    )
}

AUDIO_EXTS = {"mp3", "wav", "aac", "m4a", "flac", "ogg", "oga", "opus"}
VIDEO_EXTS = {"mp4", "mov", "m4v", "webm", "mkv", "avi"}
IMAGE_EXTS = {
    "png",
    "jpg",
    "jpeg",
    "gif",
    "webp",
    "tif",
    "tiff",
    "bmp",
    "svg",
}  # svg has no pixels

ISO_A_SIZES_MM = {
    "A0": (841.0, 1189.0),
    "A1": (594.0, 841.0),
    "A2": (420.0, 594.0),
    "A3": (297.0, 420.0),
    "A4": (210.0, 297.0),
    "A5": (148.0, 210.0),
    "A6": (105.0, 148.0),
    "A7": (74.0, 105.0),
    "A8": (52.0, 74.0),
    "A9": (37.0, 52.0),
    "A10": (26.0, 37.0),
}

# -------------------- helpers --------------------
def _hr(n: int) -> str:
    if n < 1024:
        return f"{n} Bytes"
    for u in ["KB", "MB", "GB", "TB", "PB", "EB"]:
        n /= 1024.0
        if n < 1024.0:
            return f"{n:.2f} {u}"
    return f"{n:.2f} EB"


def _format_duration(seconds: float) -> str:
    try:
        s = float(seconds)
    except Exception:
        return ""
    if s < 60:
        return f"{s:.2f} s"
    m = int(s // 60)
    sec = int(round(s - m * 60))
    if m < 60:
        return f"{m}:{sec:02d}"
    h = int(m // 60)
    m2 = int(m % 60)
    return f"{h}h {m2}m"


def _find_closest_iso_size(width_mm: float, height_mm: float) -> str:
    w, h = round(width_mm, 2), round(height_mm, 2)
    best, diff = "A4", float("inf")
    for name, (iw, ih) in ISO_A_SIZES_MM.items():
        d = min(abs(w - iw) + abs(h - ih), abs(w - ih) + abs(h - iw))
        if d < diff:
            best, diff = name, d
    return best


def _parse_s3_url(url: str) -> Tuple[Optional[str], Optional[str]]:
    p = urlparse(url)
    host, path = p.netloc, unquote(p.path.lstrip("/"))

    # virtual-hosted-style: bucket.s3.region.amazonaws.com/key or bucket.s3.amazonaws.com/key
    if host.endswith(".amazonaws.com") and (
        "s3." in host or host.startswith("s3") or "s3-" in host
    ):
        if ".s3" in host:  # virtual-hosted
            return host.split(".s3", 1)[0], path
        # path-style: s3.region.amazonaws.com/bucket/key
        parts = path.split("/", 1)
        if host.startswith("s3") and len(parts) == 2:
            return parts[0], parts[1]
    return None, None


def _s3_head(url: str) -> Optional[dict]:
    b, k = _parse_s3_url(url)
    if not b or not k:
        return None
    try:
        h = s3_client.head_object(Bucket=b, Key=k)
        return {
            "size": int(h.get("ContentLength", 0)),
            "content_type": h.get("ContentType") or DEFAULT_MIME,
            "etag": (h.get("ETag") or "").strip('"'),
            "bucket": b,
            "key": k,
        }
    except ClientError as e:
        logger.debug("S3 head failed for %s: %s", url, e)
        return None


def _guess_type(url: str) -> str:
    return mimetypes.guess_type(urlparse(url).path)[0] or DEFAULT_MIME


def _size_limit_bytes() -> Optional[int]:
    v = MAX_METADATA_BYTES
    if v in (None, 0):
        return None
    try:
        return int(v)
    except Exception:
        return None


def _ext_from_key(key: str) -> str:
    return Path(key).suffix.lower().lstrip(".")


def _resolve_lo_bin() -> Optional[str]:
    """
    Resolve LibreOffice/soffice absolute path reliably (handles names or absolute paths).
    """
    candidates = []
    if os.path.isabs(LIBREOFFICE_BIN):
        candidates.append(LIBREOFFICE_BIN)
    candidates += [LIBREOFFICE_BIN, "soffice", "libreoffice"]
    for name in candidates:
        p = shutil.which(name)
        if p:
            return p
        if os.path.isabs(name) and Path(name).exists():
            return name
    return None


# -------------------- temp + conversions --------------------
def _download_s3_to_temp(bucket: str, key: str) -> Optional[Path]:
    try:
        fd, path = tempfile.mkstemp(suffix=Path(key).suffix or ".bin")
        os.close(fd)
        with open(path, "wb") as f:
            s3_client.download_fileobj(bucket, key, f)
        return Path(path)
    except Exception as e:
        logger.debug("Download failed s3://%s/%s: %s", bucket, key, e)
        return None


def _libreoffice_pdf(src: Path) -> Optional[Path]:
    """
    Convert doc-like input to PDF via LibreOffice headless.
    Returns path to temp PDF, or None on failure.
    """
    if not DOC_PAGECOUNT_VIA_LIBREOFFICE:
        return None

    lo = _resolve_lo_bin()
    if not lo:
        logger.warning(
            "LibreOffice not found (LIBREOFFICE_BIN=%s). Skipping page-count conversion.",
            LIBREOFFICE_BIN,
        )
        return None

    try:
        outdir = Path(tempfile.mkdtemp())

        cmd = [
            lo,
            "--headless",
            "--norestore",
            "--nolockcheck",
            "--nodefault",
            "--nofirststartwizard",
            "--convert-to",
            "pdf",
            "--outdir",
            str(outdir),
            str(src),
        ]

        # Ensure LO can write config/cache even in minimal containers
        env = os.environ.copy()
        # Provide sane defaults if HOME/XDG dirs are missing
        env.setdefault("HOME", "/tmp")
        env.setdefault("XDG_CACHE_HOME", "/tmp/.cache")
        env.setdefault("XDG_CONFIG_HOME", "/tmp/.config")
        env.setdefault("TMPDIR", "/tmp")

        subprocess.run(
            cmd,
            timeout=LIBREOFFICE_TIMEOUT_SECS,
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            env=env,
        )

        pdf = outdir / (src.stem + ".pdf")

        # LO occasionally returns before fs flush; wait a moment
        for _ in range(10):
            if pdf.exists() and pdf.stat().st_size > 0:
                break
            time.sleep(0.05)

        if not pdf.exists() or pdf.stat().st_size == 0:
            logger.debug("LibreOffice produced no PDF or empty PDF for %s", src)
            return None

        return pdf
    except subprocess.TimeoutExpired:
        logger.warning(
            "LibreOffice timed out after %ss for %s", LIBREOFFICE_TIMEOUT_SECS, src
        )
        return None
    except Exception as e:
        logger.debug("LibreOffice conversion failed for %s: %s", src, e)
        return None


# -------------------- doc parsers --------------------
def _pdf_meta(p: Path) -> Dict[str, Union[str, int]]:
    out: Dict[str, Union[str, int]] = {"number_of_pages": 0, "page_size": "Unknown"}
    try:
        r = PdfReader(str(p))
        out["number_of_pages"] = len(r.pages)
        pg = r.pages[0]
        w_pt = float(pg.mediabox.width)
        h_pt = float(pg.mediabox.height)
        w_mm = (w_pt / 72.0) * 25.4
        h_mm = (h_pt / 72.0) * 25.4
        out["page_size"] = _find_closest_iso_size(w_mm, h_mm)
    except Exception as e:
        logger.debug("PDF meta failed for %s: %s", p, e)
    return out


def _pptx_meta(p: Path) -> Dict[str, Union[str, int]]:
    out: Dict[str, Union[str, int]] = {"number_of_pages": 0, "page_size": "Unknown"}
    try:
        prs = Presentation(str(p))
        slides = len(prs.slides)
        out["number_of_pages"] = slides
        out["number_of_slides"] = slides
        emu_per_inch = 914400.0
        w_in = prs.slide_width / emu_per_inch
        h_in = prs.slide_height / emu_per_inch
        out["page_size"] = _find_closest_iso_size(w_in * 25.4, h_in * 25.4)
    except Exception as e:
        logger.debug("PPTX meta failed for %s: %s", p, e)
    return out


def _docx_meta(p: Path) -> Dict[str, Union[str, int]]:
    """
    DOCX: page size via python-docx; page count via LibreOffice->PDF (enabled).
    """
    out: Dict[str, Union[str, int]] = {"number_of_pages": 0, "page_size": "Unknown"}
    try:
        doc = DocxDocument(str(p))
        if doc.sections:
            s0 = doc.sections[0]
            w_mm = float(s0.page_width.mm)
            h_mm = float(s0.page_height.mm)
            out["page_size"] = _find_closest_iso_size(w_mm, h_mm)
    except Exception as e:
        logger.debug("DOCX page size read failed for %s: %s", p, e)

    if DOCX_INCLUDE_PAGECOUNT:
        pdf = _libreoffice_pdf(p)
        if pdf:
            try:
                pdf_meta = _pdf_meta(pdf)
                n = int(pdf_meta.get("number_of_pages", 0) or 0)
                if n > 0:
                    out["number_of_pages"] = n
                if out.get("page_size") == "Unknown":
                    out["page_size"] = pdf_meta.get("page_size", "Unknown")
            finally:
                try:
                    pdf.unlink()
                except Exception:
                    pass
    return out


def _odt_meta(_p: Path) -> Dict[str, Union[str, int]]:
    # Page count normally requires conversion; handled via LO fallback if enabled.
    return {"number_of_pages": 0, "page_size": "Unknown"}


def _legacy_ppt_doc_meta(_p: Path) -> Dict[str, Union[str, int]]:
    # .ppt/.doc: rely on LO conversion below if strict/LO enabled.
    return {"number_of_pages": 0, "page_size": "Unknown"}


def _deep_doc_meta(temp_path: Path, ext_with_dot: str) -> Dict[str, Union[str, int]]:
    """
    Compute {"number_of_pages", "page_size"} for docs.
    Uses native parsers where possible; falls back to LibreOffice->PDF.
    """
    ext = (ext_with_dot or "").lower()
    meta: Dict[str, Union[str, int]] = {"number_of_pages": 0, "page_size": "Unknown"}

    try:
        if ext == ".pdf":
            return _pdf_meta(temp_path)

        if ext == ".pptx":
            return _pptx_meta(temp_path)

        if ext == ".docx":
            return _docx_meta(temp_path)

        if ext == ".odt":
            meta.update(_odt_meta(temp_path))
        elif ext in (".ppt", ".doc"):
            meta.update(_legacy_ppt_doc_meta(temp_path))

        # Fallback via LibreOffice if we still lack info and strict/LO is enabled
        need_fallback = (meta.get("number_of_pages", 0) == 0) or (
            meta.get("page_size") == "Unknown"
        )
        if need_fallback and (STRICT_DOC_PAGE_META or DOC_PAGECOUNT_VIA_LIBREOFFICE):
            pdf = _libreoffice_pdf(temp_path)
            if pdf:
                try:
                    meta.update(_pdf_meta(pdf))
                finally:
                    try:
                        pdf.unlink()
                    except Exception:
                        pass

    except Exception as e:
        logger.debug("Deep doc meta failed for %s: %s", temp_path, e)

    meta.setdefault("number_of_pages", 0)
    meta.setdefault("page_size", "Unknown")
    return meta


# -------------------- image (incremental) --------------------
def _image_dimensions_from_s3(
    bucket: str, key: str, *, max_bytes: int = 256 * 1024
) -> Optional[Tuple[int, int]]:
    """
    Incrementally fetch initial bytes and feed to PIL.ImageFile.Parser until dimensions are known,
    or max_bytes reached. SVG returns None (no pixels).
    """
    if key.lower().endswith(".svg"):
        return None

    parser = ImageFile.Parser()
    start = 0
    chunk = 64 * 1024
    while start < max_bytes:
        end = min(start + chunk - 1, max_bytes - 1)
        try:
            obj = s3_client.get_object(
                Bucket=bucket, Key=key, Range=f"bytes={start}-{end}"
            )
            data = obj["Body"].read()
            if not data:
                break
            parser.feed(data)
            if parser.image:
                return parser.image.size  # (width, height)
            start += len(data)
            chunk = min(chunk * 2, 256 * 1024)
        except ClientError as e:
            logger.debug("Range get for image failed s3://%s/%s: %s", bucket, key, e)
            break

    # small objects: single full fetch if tiny
    try:
        head = s3_client.head_object(Bucket=bucket, Key=key)
        if int(head.get("ContentLength", 0)) <= 128 * 1024:
            obj = s3_client.get_object(Bucket=bucket, Key=key)
            data = obj["Body"].read()
            parser = ImageFile.Parser()
            parser.feed(data)
            if parser.image:
                return parser.image.size
    except Exception:
        pass
    return None


# -------------------- av probe --------------------
def _ffprobe_info(url: str) -> Dict[str, Union[str, tuple]]:
    out: Dict[str, Union[str, tuple]] = {}
    try:
        timeout = int(getattr(settings, "FFPROBE_TIMEOUT_SECS", 3))
        cmd = [
            "ffprobe",
            "-v",
            "error",
            "-show_entries",
            "format=duration:stream=width,height",
            "-print_format",
            "json",
            url,
        ]
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
        if proc.returncode != 0:
            return out
        info = json.loads(proc.stdout or "{}")
        dur = info.get("format", {}).get("duration")
        if dur:
            out["duration"] = _format_duration(float(dur))
        for s in info.get("streams", []):
            if "width" in s and "height" in s:
                try:
                    out["dimensions"] = (int(s["width"]), int(s["height"]))
                    break
                except Exception:
                    pass
    except Exception:
        pass
    return out


# -------------------- main --------------------
def get_file_metadata(
    urls: List[str],
    *,
    deep_for_doc_types: bool = True,
) -> List[Dict[str, Union[str, int, float, tuple]]]:
    """
    Return list of dicts per URL:
      {
        URL, file_size, file_type,
        [number_of_pages], [page_size],
        [duration], [dimensions]
      }

    - Size/type via S3 HeadObject (fast)
    - Docs: page count + page size (LibreOffice fallback when needed; DOCX page count enabled)
    - Images: dimensions via incremental range reads (no full download)
    - Audio/Video: duration + (video) dimensions via ffprobe
    - ETag-keyed cache to avoid repeated work
    """
    out: List[Dict[str, Union[str, int, float, tuple]]] = []

    ttl = FILE_METADATA_CACHE_TTL
    size_cap = _size_limit_bytes()
    t_budget_ms = FILE_METADATA_TIME_BUDGET_MS
    started = time.monotonic()

    for url in urls:
        if not url:
            continue

        base: Dict[str, Union[str, int, tuple]] = {
            "URL": url,
            "file_size": "0 Bytes",
            "file_type": DEFAULT_MIME,
        }

        head = _s3_head(url)
        if head:
            base["file_size"] = _hr(head["size"])
            base["file_type"] = head["content_type"]
            ext = (_ext_from_key(head["key"]) or "").lower()
            cache_key = f'filemeta:{head["bucket"]}:{head["key"]}:{head["etag"]}'
            cached = cache.get(cache_key)
            if cached is not None:
                out.append({**base, **cached})
                continue

            elapsed_ms = int((time.monotonic() - started) * 1000)

            is_audio = base["file_type"].startswith("audio/") or ext in AUDIO_EXTS
            is_video = base["file_type"].startswith("video/") or ext in VIDEO_EXTS
            is_image = base["file_type"].startswith("image/") or ext in IMAGE_EXTS
            is_doc = ext in DOC_DEEP_PROBE_EXTS

            details: Dict[str, Union[str, int, tuple]] = {}

            # Images → dimensions (cheap, incremental)
            if is_image:
                dims = _image_dimensions_from_s3(head["bucket"], head["key"])
                if dims:
                    details["dimensions"] = dims

            # AV → ffprobe (respect size cap)
            if (is_audio or is_video) and (
                (size_cap is None) or (head["size"] <= size_cap)
            ):
                details.update(_ffprobe_info(url))

            # Docs → strict page meta (ignore time budget if STRICT_DOC_PAGE_META)
            if (
                deep_for_doc_types
                and is_doc
                and (
                    STRICT_DOC_PAGE_META or not t_budget_ms or elapsed_ms <= t_budget_ms
                )
            ):
                tmp = _download_s3_to_temp(head["bucket"], head["key"])
                try:
                    if tmp:
                        details.update(_deep_doc_meta(tmp, "." + ext))
                finally:
                    try:
                        if tmp:
                            tmp.unlink()
                    except Exception:
                        pass
            elif t_budget_ms and elapsed_ms > t_budget_ms:
                cache.set(cache_key, details, timeout=ttl)
                out.append({**base, **details})
                continue

            cache.set(cache_key, details, timeout=ttl)
            out.append({**base, **details})
            continue

        # Non-S3 URL fallback
        base["file_type"] = _guess_type(url)
        if base["file_type"].startswith(("audio/", "video/")):
            base.update(_ffprobe_info(url))
        out.append(base)

    return out
